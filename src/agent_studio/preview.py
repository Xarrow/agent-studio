"""文件解析 —— 把 Office 文档与思维导图解成前端能直接画的结构。

为什么放在后端
--------------
xlsx / docx / pptx / xmind 本质都是 zip + XML。要在浏览器里解，得引 SheetJS（~800KB）、
mammoth、jszip 三套库，打包体积翻倍，而且每套都有自己的坑（日期序列号、样式、字体）。
服务端已经有 openpyxl / python-docx / python-pptx，解析完输出**统一的 JSON**，
前端只负责画 —— 换前端框架不用改解析，换解析库不用改前端。

输出口径
--------
``{"kind": ..., "payload": {...}, "truncated": bool, "notes": [...]}``

* ``kind`` 是**渲染方式**而不是文件类型：spreadsheet / document / slides / outline / text
* ``truncated`` 表示"内容没展示全"，前端要明说，不能让用户以为文件就这么点
* ``notes`` 是人话补充（比如"这个 xls 是老格式，只能下载"），显示在预览顶部

所有上限都写死在模块常量里：预览的目的是"看清里面是什么"，不是"当办公软件用"。
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path
from typing import Any

#: 单个 sheet 最多画多少行 / 列（再多就该下载去 Excel 里看）
MAX_ROWS = 300
MAX_COLS = 60
#: 最多画几个 sheet
MAX_SHEETS = 12
#: 文档 / 幻灯片最多多少块
MAX_BLOCKS = 1200
#: 单元格文本上限（有的表格里塞了整篇文章）
MAX_CELL_CHARS = 800

SPREADSHEET_EXT = {".xlsx", ".xlsm"}
#: 老二进制格式：openpyxl 解不了，明确告诉用户去下载
LEGACY_SPREADSHEET_EXT = {".xls"}
DOCUMENT_EXT = {".docx"}
LEGACY_DOCUMENT_EXT = {".doc"}
SLIDES_EXT = {".pptx"}
LEGACY_SLIDES_EXT = {".ppt"}
OUTLINE_EXT = {".xmind"}

#: 这些是"能在线看内容"的类型；文件下载端点也按它放行
PARSEABLE_EXT = SPREADSHEET_EXT | DOCUMENT_EXT | SLIDES_EXT | OUTLINE_EXT


def _cell_text(v: Any) -> str:
    """单元格值 → 短文本。None 留空，日期用 ISO，浮点去掉没必要的 .0。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        # Excel 里 90% 的"数字"是浮点；不要把 3 显示成 3.0
        return str(int(v)) if v.is_integer() else repr(v)
    if hasattr(v, "isoformat"):
        try:
            # openpyxl 把 Excel 的"日期单元格"读成 datetime，哪怕它没有时间部分。
            # 直接 isoformat(sep=" ") 会显示成 "2026-10-01 00:00:00" —— 表格里全是
            # 这种零点是噪音，日期就只显示日期。
            h = getattr(v, "hour", 0)
            mi = getattr(v, "minute", 0)
            se = getattr(v, "second", 0)
            if h == 0 and mi == 0 and se == 0:
                return v.date().isoformat() if hasattr(v, "date") else v.isoformat()
            if se == 0:
                # 秒是 0 就别显示 ":00" —— 表格列宽宝贵，多三个人人都知道是零
                return v.strftime("%Y-%m-%d %H:%M")
            return v.isoformat(sep=" ")
        except Exception:  # noqa: BLE001 —— 某些第三方日期对象 isoformat 会抛
            return str(v)
    s = str(v)
    return s[:MAX_CELL_CHARS] + ("…" if len(s) > MAX_CELL_CHARS else "")


def parse_spreadsheet(path: Path) -> dict[str, Any]:
    """xlsx / xlsm → 多个 sheet 的单元格矩阵。

    用 ``data_only=True``：要的是"表里显示的值"，不是公式本身。
    副作用是**从没在 Excel 里打开过的文件**（公式没被计算过）公式位置会是空 ——
    这是 openpyxl 的已知行为，比显示一堆 ``=SUM(A1:A9)`` 更接近用户预期。
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    notes: list[str] = []
    truncated = False
    try:
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows: list[list[str]] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    truncated = True
                    break
                # 尾部全是空的列不要 —— 表的 max_column 常被远端的空单元格撑大
                cells = [_cell_text(v) for v in row[:MAX_COLS]]
                while cells and cells[-1] == "":
                    cells.pop()
                rows.append(cells)
            # 尾部连续空行也去掉（很多导出表末尾有几百个空行）
            while rows and not any(rows[-1]):
                rows.pop()
            try:
                total_rows = ws.max_row or len(rows)
                total_cols = ws.max_column or 0
            except Exception:  # noqa: BLE001 —— read_only 模式下个别文件取不到维度
                total_rows, total_cols = len(rows), 0
            sheets.append(
                {
                    "name": ws.title,
                    "rows": rows,
                    "totalRows": total_rows,
                    "totalCols": total_cols,
                }
            )
        if len(wb.worksheets) > MAX_SHEETS:
            truncated = True
            notes.append(f"共有 {len(wb.worksheets)} 个 sheet，只画了前 {MAX_SHEETS} 个")
    finally:
        wb.close()

    if truncated:
        notes.append(f"内容超出预览上限（每个 sheet 最多 {MAX_ROWS} 行 × {MAX_COLS} 列），已截断")
    return {
        "kind": "spreadsheet",
        "payload": {"sheets": sheets},
        "truncated": truncated,
        "notes": notes,
    }


def parse_document(path: Path) -> dict[str, Any]:
    """docx → 按文档顺序排的块列表（标题 / 段落 / 列表项 / 表格）。

    为什么要按 body 迭代而不是 ``doc.paragraphs`` + ``doc.tables`` 分开取：
    后者会把全文的表格全堆到最后，读起来像另一份文件。顺序错了，预览就失去意义。
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks: list[dict[str, Any]] = []
    truncated = False

    def add(block: dict[str, Any]) -> None:
        nonlocal truncated
        if len(blocks) >= MAX_BLOCKS:
            truncated = True
            return
        blocks.append(block)

    def heading_level(p: Paragraph) -> int:
        """标题级别：先看样式名（Heading 1 / 标题 1），退回 outlineLvl，再退回字重。"""
        name = (p.style.name or "") if p.style is not None else ""
        low = name.lower()
        if low in ("title", "标题"):
            return 1
        for i in range(1, 7):
            if low in (f"heading {i}", f"标题 {i}"):
                return i
        # 有些模板不用标题样式，但设了段落大纲级别
        try:
            ppr = p._p.pPr  # noqa: SLF001 —— python-docx 没有公开 API
            if ppr is not None and ppr.outlineLvl is not None:
                return min(int(ppr.outlineLvl.val) + 1, 6)
        except Exception:  # noqa: BLE001
            pass
        return 0

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            text = (p.text or "").strip()
            if not text:
                continue
            lvl = heading_level(p)
            if lvl:
                add({"t": f"h{min(lvl, 6)}", "text": text})
                continue
            name = ((p.style.name or "") if p.style is not None else "").lower()
            if "list" in name or "列表" in name:
                add({"t": "li", "text": text})
                continue
            add({"t": "p", "text": text})
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            rows: list[list[str]] = []
            for r in table.rows[:MAX_ROWS]:
                cells: list[str] = []
                for c in r.cells[:MAX_COLS]:
                    cells.append((c.text or "").strip()[:MAX_CELL_CHARS])
                rows.append(cells)
            if rows:
                add({"t": "table", "rows": rows})

    notes: list[str] = []
    if truncated:
        notes.append(f"文档很长，只展示了前 {MAX_BLOCKS} 块")
    if not blocks:
        notes.append("这个文档里没有可显示的段落或表格（可能全是图片或文本框）")
    return {"kind": "document", "payload": {"blocks": blocks}, "truncated": truncated, "notes": notes}


def parse_slides(path: Path) -> dict[str, Any]:
    """pptx → 每页的标题 / 文本 / 备注。

    不试图还原版式：幻灯片的价值在"讲了哪几件事"，字号颜色对预览没用，
    而排版还原是所有 Office 预览里最容易做砸的一环。这里老实出文本，
    并按页分组 —— 用户一眼能看出结构。
    """
    from pptx import Presentation

    prs = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    for idx, slide in enumerate(prs.slides, start=1):
        if idx > MAX_BLOCKS:
            break
        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:  # noqa: BLE001
            title = ""
        lines: list[str] = []
        for shape in slide.shapes:
            try:
                if not shape.has_text_frame:
                    continue
                txt = (shape.text_frame.text or "").strip()
                if not txt:
                    continue
                if title and txt == title:
                    continue  # 标题已经在头顶显示了，不要在正文重复一遍
                lines.append(txt)
            except Exception:  # noqa: BLE001 —— 个别图形没有 text_frame
                continue
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:  # noqa: BLE001
            notes = ""
        slides.append({"n": idx, "title": title, "lines": lines, "notes": notes})

    truncated = len(prs.slides) > len(slides)
    out_notes: list[str] = []
    if truncated:
        out_notes.append(f"共 {len(prs.slides)} 页，只展示了前 {len(slides)} 页")
    if not slides:
        out_notes.append("这份幻灯片没有可读的文字（可能每页都是整张图片）")
    return {"kind": "slides", "payload": {"slides": slides}, "truncated": truncated, "notes": out_notes}


def _topic_from_json(t: dict[str, Any], depth: int = 0) -> dict[str, Any]:
    """XMind 2020+ 的 content.json 主题 → {title, children, note}。"""
    title = str(t.get("title") or "").strip()
    label = str(t.get("label") or "").strip()
    if not title and label:
        title = label
    kids: list[dict[str, Any]] = []
    children = t.get("children") or {}
    attached = children.get("attached") if isinstance(children, dict) else None
    for c in attached or []:
        if isinstance(c, dict) and depth < 24:
            kids.append(_topic_from_json(c, depth + 1))
    notes = t.get("notes")
    note = ""
    if isinstance(notes, dict):
        plain = notes.get("plain")
        if isinstance(plain, dict):
            note = str(plain.get("content") or "").strip()
    node: dict[str, Any] = {"title": title}
    if kids:
        node["children"] = kids
    if note:
        node["note"] = note
    return node


def parse_xmind(path: Path) -> dict[str, Any]:
    """xmind → 多张画布的树。

    xmind 就是 zip：新版（XMind 2020 / 8+）是 ``content.json``，老版是 ``content.xml``。
    两种都处理 —— 很多人手里的 .xmind 还是老版本导出的。
    """
    sheets: list[dict[str, Any]] = []
    notes: list[str] = []

    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        if "content.json" in names:
            raw = json.loads(z.read("content.json").decode("utf-8", "replace"))
            for item in raw if isinstance(raw, list) else [raw]:
                if not isinstance(item, dict):
                    continue
                root = item.get("rootTopic") or {}
                sheets.append(
                    {
                        "title": str(item.get("title") or "画布"),
                        "root": _topic_from_json(root) if isinstance(root, dict) else {"title": ""},
                    }
                )
        elif "content.xml" in names:
            import xml.etree.ElementTree as ET

            tree = ET.fromstring(z.read("content.xml"))
            for sheet in tree.iter():
                if not sheet.tag.endswith("sheet"):
                    continue
                stitle = "画布"
                root_node: dict[str, Any] = {"title": ""}
                for ch in sheet:
                    if ch.tag.endswith("title"):
                        stitle = (ch.text or "").strip() or stitle
                    elif ch.tag.endswith("topic"):
                        root_node = _topic_from_xml(ch, 0)
                sheets.append({"title": stitle, "root": root_node})
        else:
            raise ValueError("这个 xmind 里没有 content.json / content.xml")

    if not sheets:
        notes.append("没解析出内容（文件可能损坏或格式特殊）")
    return {"kind": "outline", "payload": {"sheets": sheets}, "truncated": False, "notes": notes}


def _topic_from_xml(el: Any, depth: int) -> dict[str, Any]:
    title = ""
    kids: list[dict[str, Any]] = []
    for ch in el:
        if ch.tag.endswith("title") and not title:
            title = (ch.text or "").strip()
        elif ch.tag.endswith("children") and depth < 24:
            for topics in ch:
                for topic in topics:
                    if topic.tag.endswith("topic"):
                        kids.append(_topic_from_xml(topic, depth + 1))
    node: dict[str, Any] = {"title": title}
    if kids:
        node["children"] = kids
    return node


def parse(path: Path) -> dict[str, Any]:
    """按后缀分派。解不了的格式抛 ValueError，由调用方转成人话提示。"""
    ext = path.suffix.lower()
    size = path.stat().st_size
    head = {
        "name": path.name,
        "size": size,
        "ext": ext.lstrip("."),
    }

    if ext in LEGACY_SPREADSHEET_EXT | LEGACY_DOCUMENT_EXT | LEGACY_SLIDES_EXT:
        raise ValueError(
            f"{ext} 是 Office 97-2003 的老格式（二进制），在线解不了 —— "
            f"在 Office 里另存为 {ext}x 就能预览"
        )
    if ext in SPREADSHEET_EXT:
        return {**head, **parse_spreadsheet(path)}
    if ext in DOCUMENT_EXT:
        return {**head, **parse_document(path)}
    if ext in SLIDES_EXT:
        return {**head, **parse_slides(path)}
    if ext in OUTLINE_EXT:
        return {**head, **parse_xmind(path)}
    raise ValueError(f"不支持预览的类型：{ext}")
